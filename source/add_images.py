from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from pathlib import Path
import shutil
ids=["exec-658b55c5-9453-4ee6-904a-a21ecd63de5b.png","exec-d34bf3fc-c574-4f85-8af8-881a23f2a4e5.png","exec-d6271ca9-9f18-463e-b626-9d94f41983c5.png","exec-51e58fdf-6682-48a5-99c0-d6066a336907.png","exec-1983c05e-0b9f-43ac-a35d-044e36318934.png","exec-724979e5-fb23-4e6e-967b-b6e162c4b352.png","exec-941107e4-3143-4290-9d4b-d1563ac8a3cf.png","exec-49c6d000-53b1-47c8-a482-41688c903411.png","exec-5d7fbf15-7378-488b-b887-e0835d5e7da3.png","exec-edbfce48-bb16-4bff-ad75-37fdf302a969.png","exec-b20dbc21-f02e-4602-b026-2c458885dbb9.png","exec-cf4fe0c4-fc25-46cd-98f2-debb5c29232b.png","exec-45f7ede1-a78d-4e43-89e5-b850fa3799c1.png","exec-b20a7e06-7ecf-48e7-b913-34e58e0ffcdb.png","exec-5cf0b452-ccc6-4df0-a645-5e4cbe00df88.png","exec-31a9d28f-1a03-4bf2-bdb0-caedf0ea003b.png","exec-fdc5a988-ab4f-45f7-9836-d4d8260d1746.png","exec-46dc1234-ad16-490b-94e1-4d1c18dc4e6b.png","exec-9f13b0e5-a114-4608-8f0c-e2e4a81e29c0.png","exec-4e407cfd-7080-4d6d-a144-ef66629c67fa.png"]
src=Path('C:/Users/surak/.codex/generated_images/01a18ca2-b450-74e3-9e8c-1bf9376e0114')
assets=Path('D:/www/cook-book/food_images'); assets.mkdir(exist_ok=True)
doc=Document('D:/www/cook-book/หนังสือสูตรอาหารไทย_20เมนู.docx')
heads=[p for p in doc.paragraphs if p.style.name=='Heading 1' and len(p.text)>3 and p.text[:2].isdigit() and p.text[2:4]=='  ']
assert len(heads)==20
for i,(heading,fn) in enumerate(zip(heads,ids),1):
 dest=assets / f'{i:02d}.png'; shutil.copyfile(src / fn,dest)
 node=OxmlElement('w:p'); heading._p.addnext(node); para=Paragraph(node,heading._parent)
 para.paragraph_format.space_after=Pt(4)
 para.add_run().add_picture(str(dest),width=Cm(4))
doc.save('D:/www/cook-book/หนังสือสูตรอาหารไทย_20เมนู_พร้อมภาพ.docx')
print('recipes',len(heads),'images',len(doc.inline_shapes))

